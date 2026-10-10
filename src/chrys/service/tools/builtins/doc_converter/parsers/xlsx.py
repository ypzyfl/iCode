# Copyright (c) Huawei Technologies Co., Ltd. 2026. All rights reserved.

"""XLSX parser — converts Excel spreadsheets to Markdown using openpyxl."""

from __future__ import annotations

from typing import Any

from chrys.service.tools.builtins.doc_converter.parsers._table import rows_to_markdown_table
from chrys.service.tools.builtins.doc_converter.parsers.base import DocumentImageSink, ParsedDocument


def _load_worksheets(path: str) -> Any:
    """Open *path* read-only with its worksheets and without its chart sheets.

    A chart sheet holds no cells, and openpyxl loads its drawings even
    read-only: every picture on one goes through Pillow with every decoder.
    """
    from openpyxl.reader.excel import ExcelReader

    class _WorksheetReader(ExcelReader):
        def read_chartsheet(self, sheet: object, rel: object) -> None:
            """Leave the chart sheet, and the drawings it would load, unread."""

    reader = _WorksheetReader(path, read_only=True, data_only=True)
    reader.read()
    return reader.wb


class XlsxParser:
    """Convert XLSX files to Markdown tables via ``openpyxl``."""

    @property
    def supported_extensions(self) -> frozenset[str]:
        return frozenset({".xlsx"})

    def parse(self, path: str, *, image_sink: DocumentImageSink | None = None) -> ParsedDocument:
        wb = _load_worksheets(path)
        parts: list[str] = []
        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                parts.append(f"# Sheet: {sheet_name}")
                if not rows:
                    parts.append("(empty sheet)")
                    continue
                parts.append(rows_to_markdown_table(rows))
        finally:
            wb.close()
        return ParsedDocument(markdown="\n\n".join(parts))
