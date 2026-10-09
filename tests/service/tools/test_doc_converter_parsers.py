# Copyright (c) Huawei Technologies Co., Ltd. 2026. All rights reserved.

"""Tests for the doc_converter parser registry and the pdf/docx/pptx/xlsx/xls/_table parsers."""

from __future__ import annotations

import ast
import zipfile
from io import BytesIO
from pathlib import Path
from typing import IO
from unittest.mock import MagicMock, patch

import pytest

from chrys.service.tools.builtins.doc_converter.artifacts import (
    MAX_IMAGE_OCCURRENCES,
    DocumentImageSink,
)
from chrys.service.tools.builtins.doc_converter.parsers.base import DOCUMENT_IMAGE_DECODE_FORMATS
from tests.service.tools._doc_converter_fakes import (
    _png_bytes,
    _session_artifact_path,
)
from tests.support.images import image_bytes


class _FakePdfImages:
    """Stand-in for pypdf's ``page.images``: ``keys()`` lists every candidate, ``[]`` builds one image object per lookup."""

    def __init__(self, entries: dict[str, dict[str, object]]) -> None:
        self._entries = entries
        self.keys_materialized = 0
        self.retrieval_count = 0

    def keys(self) -> list[str]:
        self.keys_materialized += len(self._entries)
        return list(self._entries)

    def __getitem__(self, key: str) -> MagicMock:
        self.retrieval_count += 1
        return MagicMock(**self._entries[key])


def _pdf_image(
    data: bytes, name: str, *, is_displayed: bool = True, indirect_reference: object = None
) -> dict[str, object]:
    return {"data": data, "name": name, "is_displayed": is_displayed, "indirect_reference": indirect_reference}


@pytest.fixture
def mocked_pypdf_images(monkeypatch: pytest.MonkeyPatch) -> None:
    """Tests that replace pypdf itself have no pypdf image module to restrict."""
    from chrys.service.tools.builtins.doc_converter.parsers import pdf

    def leave_unrestricted() -> None:
        return None

    monkeypatch.setattr(pdf, "_restrict_pypdf_image_decoders", leave_unrestricted)


class _PillowOpenRecorder:
    """Wraps the real ``PIL.Image.open``, recording the formats each call allows and the format it decoded."""

    def __init__(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from PIL import Image

        self._open = Image.open
        self.allowed: list[tuple[str, ...] | None] = []
        self.decoded: list[str | None] = []
        monkeypatch.setattr(Image, "open", self.open)

    def open(self, fp: str | IO[bytes], mode: str = "r", formats: list[str] | tuple[str, ...] | None = None) -> object:
        self.allowed.append(None if formats is None else tuple(formats))
        image = self._open(fp, mode, formats)
        self.decoded.append(image.format)
        return image


def _pdf_with_image(filter_name: str, stream: bytes) -> bytes:
    """A one-page PDF drawing one 2-by-2 RGB image XObject whose stream *filter_name* claims to encode."""
    content = b"q 20 0 0 20 0 0 cm /Im0 Do Q"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 100 100] "
            b"/Resources << /XObject << /Im0 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceRGB /BitsPerComponent 8 "
        b"/Filter /%s /Length %d >>\nstream\n%s\nendstream" % (filter_name.encode(), len(stream), stream),
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


# ---------------------------------------------------------------------------
# Parser registry
# ---------------------------------------------------------------------------


def test_parser_registry_all_extensions() -> None:
    """All expected extensions are registered."""
    from chrys.service.tools.builtins.doc_converter.registry import supported_extensions

    exts = supported_extensions()
    assert ".pdf" in exts
    assert ".docx" in exts
    assert ".pptx" in exts
    assert ".xlsx" in exts
    assert ".xls" in exts
    assert ".epub" not in exts


def test_parser_registry_get_parser() -> None:
    """get_parser returns correct parser type for each extension."""
    from chrys.service.tools.builtins.doc_converter.parsers.docx import DocxParser
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser
    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser
    from chrys.service.tools.builtins.doc_converter.parsers.xls import XlsParser
    from chrys.service.tools.builtins.doc_converter.parsers.xlsx import XlsxParser
    from chrys.service.tools.builtins.doc_converter.registry import get_parser

    assert isinstance(get_parser(".pdf"), PdfParser)
    assert isinstance(get_parser(".docx"), DocxParser)
    assert isinstance(get_parser(".pptx"), PptxParser)
    assert isinstance(get_parser(".xlsx"), XlsxParser)
    assert isinstance(get_parser(".xls"), XlsParser)
    assert get_parser(".epub") is None
    assert get_parser(".txt") is None


def test_parser_registry_case_insensitive() -> None:
    from chrys.service.tools.builtins.doc_converter.registry import get_parser

    assert get_parser(".PDF") is not None
    assert get_parser(".Docx") is not None


# ---------------------------------------------------------------------------
# Individual parser unit tests (mock the library imports)
# ---------------------------------------------------------------------------


def test_pdf_parser_output() -> None:
    """PdfParser produces page-based Markdown headings."""
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    mock_page1 = MagicMock()
    mock_page1.extract_text.return_value = "Hello world"
    mock_page2 = MagicMock()
    mock_page2.extract_text.return_value = "Second page content"

    mock_reader = MagicMock()
    mock_reader.pages = [mock_page1, mock_page2]

    mock_pypdf = MagicMock()
    mock_pypdf.PdfReader.return_value = mock_reader

    with patch.dict("sys.modules", {"pypdf": mock_pypdf}):
        parser = PdfParser()
        result = parser.parse("/fake.pdf")

    assert "# Page 1" in result.markdown
    assert "Hello world" in result.markdown
    assert "# Page 2" in result.markdown
    assert "Second page content" in result.markdown


def test_pdf_parser_extracts_image_only_page_and_keeps_no_text_placeholder(tmp_path: Path) -> None:
    from PIL import Image

    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    pdf = tmp_path / "image.pdf"
    Image.new("RGB", (8, 8), (255, 0, 0)).save(pdf, format="PDF")
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="image")

    result = PdfParser().parse(str(pdf), image_sink=sink)

    assert "# Page 1\n\n(no text content)" in result.markdown
    assert len(result.visuals) == 1
    assert _session_artifact_path(tmp_path / "session", result.visuals[0].reference).exists()


def test_pdf_parser_extracts_a_dct_image_through_the_restricted_decoders(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    pillow = _PillowOpenRecorder(monkeypatch)
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(_pdf_with_image("DCTDecode", image_bytes("JPEG")))
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="scan")

    result = PdfParser().parse(str(pdf), image_sink=sink)

    assert len(result.visuals) == 1
    assert result.warnings == ()
    assert "JPEG" in pillow.decoded
    assert all(formats is not None and set(formats) <= set(DOCUMENT_IMAGE_DECODE_FORMATS) for formats in pillow.allowed)


@pytest.mark.parametrize("image_format", ["PCX", "EPS"])
def test_pdf_parser_never_decodes_a_foreign_format_behind_dct(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, image_format: str
) -> None:
    """pypdf opens a /DCTDecode stream with every Pillow decoder (EPS runs Ghostscript) unless restricted."""
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    pillow = _PillowOpenRecorder(monkeypatch)
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(_pdf_with_image("DCTDecode", image_bytes(image_format)))
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="scan")

    result = PdfParser().parse(str(pdf), image_sink=sink)

    assert pillow.allowed
    assert all(formats is not None and set(formats) <= set(DOCUMENT_IMAGE_DECODE_FORMATS) for formats in pillow.allowed)
    assert image_format not in pillow.decoded
    assert result.visuals == ()
    assert len(result.warnings) == 1
    assert "could not be decoded or normalized" in result.warnings[0]


def test_pdf_image_restriction_is_installed_once(monkeypatch: pytest.MonkeyPatch) -> None:
    from PIL import Image
    from pypdf.generic import _image_xobject

    from chrys.service.tools.builtins.doc_converter.parsers.pdf import _restrict_pypdf_image_decoders

    monkeypatch.setattr(_image_xobject, "Image", Image)
    _restrict_pypdf_image_decoders()
    installed = _image_xobject.Image
    _restrict_pypdf_image_decoders()

    assert installed is not Image
    assert _image_xobject.Image is installed


def test_pdf_parser_extracts_a_jbig2_image_that_jbig2dec_wrote_as_pbm(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """jbig2dec built without libpng writes PBM even when pypdf asks it for PNG."""
    from pypdf import filters

    from chrys.service.tools.builtins.doc_converter.artifacts import DocumentImageSink
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    def jbig2dec_writing_pbm(data: bytes, decode_parms: object = None) -> bytes:
        return b"P4\n2 2\n\x80\x40"

    monkeypatch.setattr(filters.JBIG2Decode, "decode", staticmethod(jbig2dec_writing_pbm))
    pillow = _PillowOpenRecorder(monkeypatch)
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(_pdf_with_image("JBIG2Decode", b"jbig2 segments"))
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="scan")

    result = PdfParser().parse(str(pdf), image_sink=sink)

    assert len(result.visuals) == 1
    assert result.warnings == ()
    assert "PPM" in pillow.decoded
    assert "PPM" not in DOCUMENT_IMAGE_DECODE_FORMATS


def test_pdf_image_restriction_keeps_only_reviewed_formats_pypdf_names() -> None:
    from PIL import Image, UnidentifiedImageError

    from chrys.service.tools.builtins.doc_converter.parsers.pdf import _DocumentFormatsImage

    restricted = _DocumentFormatsImage(Image)
    pcx = image_bytes("PCX")

    assert Image.open(BytesIO(pcx), formats=("PNG", "PCX")).format == "PCX"
    with pytest.raises(UnidentifiedImageError):
        restricted.open(BytesIO(pcx), formats=("PNG", "PCX"))
    assert restricted.open(BytesIO(image_bytes("PPM")), formats=("PNG", "PPM")).format == "PPM"


def test_pdf_image_restriction_keeps_every_format_the_installed_pypdf_names() -> None:
    """A pypdf upgrade that names a new format fails here instead of skipping those images: review it first."""
    from pypdf.generic import _image_xobject

    from chrys.service.tools.builtins.doc_converter.parsers.pdf import _PYPDF_NAMED_FORMATS

    tree = ast.parse(Path(str(_image_xobject.__file__)).read_text(encoding="utf-8"))
    named = [
        keyword.value
        for node in ast.walk(tree)
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr == "open"
        for keyword in node.keywords
        if keyword.arg == "formats"
    ]

    assert named
    assert {name for value in named for name in ast.literal_eval(value)} <= set(_PYPDF_NAMED_FORMATS)


def test_pdf_image_restriction_refuses_a_pypdf_that_stopped_using_pil_image(monkeypatch: pytest.MonkeyPatch) -> None:
    from pypdf.generic import _image_xobject

    from chrys.service.tools.builtins.doc_converter.parsers.pdf import _restrict_pypdf_image_decoders

    replacement = object()
    monkeypatch.setattr(_image_xobject, "Image", replacement)

    with pytest.raises(RuntimeError, match=r"pypdf no longer decodes images through PIL\.Image"):
        _restrict_pypdf_image_decoders()
    assert _image_xobject.Image is replacement


@pytest.mark.usefixtures("mocked_pypdf_images")
def test_pdf_parser_deduplicates_repeated_indirect_reference(tmp_path: Path) -> None:
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    image_data = _png_bytes()
    indirect_reference = object()

    pages = []
    for _ in range(2):
        page = MagicMock()
        page.extract_text.return_value = "text"
        page.images = _FakePdfImages(
            {"/Im0": _pdf_image(image_data, "logo.png", indirect_reference=indirect_reference)}
        )
        pages.append(page)
    mock_pypdf = MagicMock()
    mock_pypdf.PdfReader.return_value.pages = pages
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="report")

    with (
        patch.dict("sys.modules", {"pypdf": mock_pypdf}),
        patch.object(sink, "save_image", wraps=sink.save_image) as save_image,
    ):
        result = PdfParser().parse("/fake.pdf", image_sink=sink)

    assert len(result.visuals) == 2
    assert result.visuals[0].reference == result.visuals[1].reference
    assert save_image.call_count == 1
    assert len(list((tmp_path / "session" / "doc_converter").iterdir())) == 1


@pytest.mark.usefixtures("mocked_pypdf_images")
def test_pdf_parser_skips_undisplayed_resource_without_writing(tmp_path: Path) -> None:
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    page = MagicMock()
    page.extract_text.return_value = "text"
    page.images = _FakePdfImages(
        {"/Unused": _pdf_image(_png_bytes(), "unused.png", is_displayed=False, indirect_reference=object())}
    )
    mock_pypdf = MagicMock()
    mock_pypdf.PdfReader.return_value.pages = [page]
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="report")

    with patch.dict("sys.modules", {"pypdf": mock_pypdf}):
        result = PdfParser().parse("/fake.pdf", image_sink=sink)

    assert result.visuals == ()
    assert not (tmp_path / "session" / "doc_converter").exists()


@pytest.mark.usefixtures("mocked_pypdf_images")
def test_pdf_parser_corrupt_image_preserves_text_and_other_images(tmp_path: Path) -> None:
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    page = MagicMock()
    page.extract_text.return_value = "Preserved page text"
    page.images = _FakePdfImages(
        {
            "/Corrupt": _pdf_image(b"not-an-image", "/Corrupt.png"),
            "/Valid": _pdf_image(_png_bytes(), "/Valid.png"),
        }
    )
    mock_pypdf = MagicMock()
    mock_pypdf.PdfReader.return_value.pages = [page]
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="report")

    with patch.dict("sys.modules", {"pypdf": mock_pypdf}):
        result = PdfParser().parse("/fake.pdf", image_sink=sink)

    assert "Preserved page text" in result.markdown
    assert len(result.visuals) == 1
    assert len(result.warnings) == 1
    assert "could not be decoded or normalized" in result.warnings[0]


@pytest.mark.usefixtures("mocked_pypdf_images")
@pytest.mark.parametrize(
    ("key_template", "image_name"),
    [("/Im{index}", "logo.png"), ("~{index}~", "inline.png")],
    ids=["xobject", "inline"],
)
def test_pdf_parser_caps_image_retrievals_after_materializing_keys(
    tmp_path: Path, key_template: str, image_name: str
) -> None:
    """``keys()`` materializes every candidate (pypdf scans inline images eagerly); retrievals stop at the cap."""
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser

    candidate_count = MAX_IMAGE_OCCURRENCES + 2
    images = _FakePdfImages(
        {key_template.format(index=index): _pdf_image(_png_bytes(), image_name) for index in range(candidate_count)}
    )
    page = MagicMock()
    page.extract_text.return_value = "text"
    page.images = images
    mock_pypdf = MagicMock()
    mock_pypdf.PdfReader.return_value.pages = [page]
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="report")

    with patch.dict("sys.modules", {"pypdf": mock_pypdf}):
        result = PdfParser().parse("/fake.pdf", image_sink=sink)

    assert images.keys_materialized == candidate_count
    assert images.retrieval_count == MAX_IMAGE_OCCURRENCES
    assert len(result.visuals) == MAX_IMAGE_OCCURRENCES
    assert len(result.warnings) == 1
    assert "2 image candidate(s)" in result.warnings[0]
    assert sink.warnings == result.warnings


def test_docx_parser_headings_and_tables() -> None:
    """DocxParser maps heading styles to Markdown headings and extracts tables in order."""
    from chrys.service.tools.builtins.doc_converter.parsers.docx import DocxParser

    # Real classes so isinstance() checks work with the mocked docx.table.Table
    class _FakeDocxTable:
        def __init__(self, row_data: list[list[str]]):
            self.rows = [MagicMock(cells=[MagicMock(text=c) for c in r]) for r in row_data]

    mock_para1 = MagicMock()
    mock_para1.text = "My Title"
    mock_para1.style.name = "Heading 1"

    mock_para2 = MagicMock()
    mock_para2.text = "Some body text"
    mock_para2.style.name = "Normal"

    mock_table = _FakeDocxTable([["Name", "Age"], ["Alice", "30"]])

    mock_para3 = MagicMock()
    mock_para3.text = "Subsection"
    mock_para3.style.name = "Heading 2"

    mock_doc = MagicMock()
    mock_doc.iter_inner_content.return_value = iter([mock_para1, mock_para2, mock_table, mock_para3])

    mock_docx = MagicMock()
    mock_docx.Document.return_value = mock_doc
    mock_table_mod = MagicMock()
    mock_table_mod.Table = _FakeDocxTable

    with patch.dict("sys.modules", {"docx": mock_docx, "docx.table": mock_table_mod}):
        parser = DocxParser()
        result = parser.parse("/fake.docx")

    assert "# My Title" in result.markdown
    assert "Some body text" in result.markdown
    assert "| Name | Age |" in result.markdown
    assert "| Alice | 30 |" in result.markdown
    assert "## Subsection" in result.markdown
    # Verify ordering: title before table before subsection
    assert result.markdown.index("My Title") < result.markdown.index("Name | Age") < result.markdown.index("Subsection")


def test_docx_parser_extracts_package_wide_header_image(tmp_path: Path) -> None:
    from docx import Document

    from chrys.service.tools.builtins.doc_converter.parsers.docx import DocxParser

    document = Document()
    document.add_paragraph("Body text")
    document.sections[0].header.paragraphs[0].add_run().add_picture(BytesIO(_png_bytes()))
    path = tmp_path / "header.docx"
    document.save(path)
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="header")

    result = DocxParser().parse(str(path), image_sink=sink)

    assert result.markdown == "Body text"
    assert len(result.visuals) == 1
    assert result.visuals[0].location == "Document"
    assert _session_artifact_path(tmp_path / "session", result.visuals[0].reference).exists()


def test_docx_parser_package_image_parts_deduplicate_reused_blob(tmp_path: Path) -> None:
    from docx import Document

    from chrys.service.tools.builtins.doc_converter.parsers.docx import DocxParser

    image = BytesIO(_png_bytes())
    document = Document()
    document.add_picture(image)
    image.seek(0)
    document.add_picture(image)
    path = tmp_path / "reused.docx"
    document.save(path)
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="reused")

    result = DocxParser().parse(str(path), image_sink=sink)

    assert result.markdown == ""
    assert len(result.visuals) == 1
    assert len(list((tmp_path / "session" / "doc_converter").iterdir())) == 1


def test_pptx_parser_slides() -> None:
    """PptxParser produces slide-based Markdown headings."""
    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    mock_title = MagicMock()
    mock_title.text = "Intro Slide"

    mock_shape = MagicMock()
    mock_shape.has_text_frame = True
    mock_shape.has_table = False
    mock_para = MagicMock()
    mock_para.text = "Bullet point"
    mock_shape.text_frame.paragraphs = [mock_para]

    mock_slide = MagicMock()
    mock_slide.shapes.title = mock_title
    mock_slide.shapes.__iter__ = MagicMock(return_value=iter([mock_shape]))

    mock_prs = MagicMock()
    mock_prs.slides = [mock_slide]

    mock_pptx = MagicMock()
    mock_pptx.Presentation.return_value = mock_prs

    with patch.dict("sys.modules", {"pptx": mock_pptx}):
        parser = PptxParser()
        result = parser.parse("/fake.pptx")

    assert "# Slide 1: Intro Slide" in result.markdown
    assert "Bullet point" in result.markdown


def test_pptx_parser_table_shape() -> None:
    """PptxParser extracts tables from table shapes."""
    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    # Text shape
    text_shape = MagicMock()
    text_shape.has_table = False
    text_shape.has_text_frame = True
    text_shape.text_frame.paragraphs = [MagicMock(text="Intro text")]

    # Table shape
    table_shape = MagicMock()
    table_shape.has_table = True
    table_shape.has_text_frame = False
    table_shape.table.rows = [
        MagicMock(cells=[MagicMock(text="Col A"), MagicMock(text="Col B")]),
        MagicMock(cells=[MagicMock(text="val1"), MagicMock(text="val2")]),
    ]

    mock_slide = MagicMock()
    mock_slide.shapes.title = None
    mock_slide.shapes.__iter__ = MagicMock(return_value=iter([text_shape, table_shape]))

    mock_prs = MagicMock()
    mock_prs.slides = [mock_slide]

    mock_pptx = MagicMock()
    mock_pptx.Presentation.return_value = mock_prs

    with patch.dict("sys.modules", {"pptx": mock_pptx}):
        parser = PptxParser()
        result = parser.parse("/fake.pptx")

    assert "# Slide 1" in result.markdown
    assert "Intro text" in result.markdown
    assert "| Col A | Col B |" in result.markdown
    assert "| val1 | val2 |" in result.markdown


def test_pptx_parser_extracts_top_level_and_grouped_pictures(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    image = tmp_path / "pixel.png"
    image.write_bytes(_png_bytes())
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    top_level = slide.shapes.add_picture(str(image), Inches(1), Inches(1))
    slide.shapes.add_group_shape([top_level])
    deck = tmp_path / "deck.pptx"
    presentation.save(deck)
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="deck")

    result = PptxParser().parse(str(deck), image_sink=sink)

    assert len(result.visuals) == 1
    assert result.visuals[0].location == "Slide 1"
    assert _session_artifact_path(tmp_path / "session", result.visuals[0].reference).exists()


def test_pptx_parser_extracts_populated_picture_placeholder(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.shapes.picture import Picture

    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[8])
    placeholder = next(
        candidate for candidate in slide.placeholders if candidate.placeholder_format.type == PP_PLACEHOLDER.PICTURE
    )
    populated = placeholder.insert_picture(BytesIO(_png_bytes()))
    assert isinstance(populated, Picture)
    assert populated.shape_type.name == "PLACEHOLDER"
    deck = tmp_path / "placeholder.pptx"
    presentation.save(deck)
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="placeholder")

    result = PptxParser().parse(str(deck), image_sink=sink)

    assert len(result.visuals) == 1
    assert _session_artifact_path(tmp_path / "session", result.visuals[0].reference).exists()


def test_pptx_parser_repeated_logo_across_slides_stores_once(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    image = tmp_path / "logo.png"
    image.write_bytes(_png_bytes())
    presentation = Presentation()
    for _ in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(image), Inches(1), Inches(1))
    deck = tmp_path / "repeated-logo.pptx"
    presentation.save(deck)
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="repeated-logo")

    result = PptxParser().parse(str(deck), image_sink=sink)

    assert len(result.visuals) == 2
    assert result.visuals[0].reference == result.visuals[1].reference
    assert len(list((tmp_path / "session" / "doc_converter").iterdir())) == 1


def test_pptx_parser_skips_linked_only_picture_with_bounded_warning(tmp_path: Path) -> None:
    from pptx.shapes.picture import Picture

    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser

    class _LinkedPicture(Picture):
        @property
        def has_table(self) -> bool:
            return False

        @property
        def has_text_frame(self) -> bool:
            return False

        @property
        def image(self):
            raise ValueError("no embedded image")

    linked = object.__new__(_LinkedPicture)
    slide = MagicMock()
    slide.shapes.title = None
    slide.shapes.__iter__ = MagicMock(side_effect=lambda: iter([linked]))
    presentation = MagicMock()
    presentation.slides = [slide]
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="linked")

    with patch("pptx.Presentation", return_value=presentation):
        result = PptxParser().parse("/fake.pptx", image_sink=sink)

    assert result.visuals == ()
    assert len(result.warnings) == 1
    assert "linked image(s)" in result.warnings[0]


def test_xlsx_parser_tables(tmp_path: Path) -> None:
    """XlsxParser produces sheet-based Markdown tables."""
    from openpyxl import Workbook

    from chrys.service.tools.builtins.doc_converter.parsers.xlsx import XlsxParser

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Sheet1"
    for row in [("Name", "Age"), ("Alice", 30), ("Bob", 25)]:
        worksheet.append(row)
    spreadsheet = tmp_path / "people.xlsx"
    workbook.save(spreadsheet)
    workbook.close()

    result = XlsxParser().parse(str(spreadsheet), image_sink=MagicMock(spec=DocumentImageSink))

    assert "# Sheet: Sheet1" in result.markdown
    assert "| Name | Age |" in result.markdown
    assert "| Alice | 30 |" in result.markdown
    assert result.warnings == ()


def _workbook_with_picture_chart_sheet(tmp_path: Path, picture: bytes) -> Path:
    """Save a "Data" worksheet and a "Chart" sheet, then make "Chart" a chart sheet showing *picture*.

    openpyxl writes pictures on worksheets only, so the second sheet is saved as
    a worksheet with one picture and its parts are rewritten into a chart sheet.
    """
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SpreadsheetImage

    picture_path = tmp_path / "placeholder.png"
    picture_path.write_bytes(_png_bytes())
    workbook = Workbook()
    workbook.active.title = "Data"
    workbook.active.append(("Region", "Sales"))
    workbook.active.append(("North", 12))
    chart = workbook.create_sheet("Chart")
    chart.add_image(SpreadsheetImage(picture_path), "A1")
    saved = tmp_path / "saved.xlsx"
    workbook.save(saved)
    workbook.close()

    worksheet_type = b"/relationships/worksheet"
    rewritten = tmp_path / "report.xlsx"
    with zipfile.ZipFile(saved) as source, zipfile.ZipFile(rewritten, "w") as target:
        workbook_rels = source.read("xl/_rels/workbook.xml.rels")
        sheet2 = b'Target="/xl/worksheets/sheet2.xml"'
        relationships = workbook_rels.split(b"<Relationship ")
        workbook_rels = b"<Relationship ".join(
            part.replace(worksheet_type, b"/relationships/chartsheet") if sheet2 in part else part
            for part in relationships
        )
        parts = {
            "xl/_rels/workbook.xml.rels": workbook_rels,
            "xl/worksheets/sheet2.xml": (
                b'<chartsheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                b'<sheetViews><sheetView workbookViewId="0"/></sheetViews><drawing r:id="rId1"/></chartsheet>'
            ),
            "xl/media/image1.png": picture,
        }
        for item in source.infolist():
            target.writestr(item, parts.get(item.filename, source.read(item.filename)))
    return rewritten


def test_xlsx_parser_converts_worksheets_and_leaves_chart_sheet_pictures_undecoded(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """openpyxl loads a chart sheet's drawings even read-only, opening each picture with every decoder."""
    from chrys.service.tools.builtins.doc_converter.parsers.xlsx import XlsxParser

    spreadsheet = _workbook_with_picture_chart_sheet(tmp_path, image_bytes("PCX"))
    pillow = _PillowOpenRecorder(monkeypatch)

    result = XlsxParser().parse(str(spreadsheet))

    assert pillow.allowed == []
    assert "# Sheet: Data" in result.markdown
    assert "| North | 12 |" in result.markdown
    assert "Chart" not in result.markdown


def test_xlsx_parser_with_embedded_image_stays_text_only_without_result_warning(tmp_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SpreadsheetImage

    from chrys.service.tools.builtins.doc_converter.parsers.xlsx import XlsxParser

    image_path = tmp_path / "chart.png"
    image_path.write_bytes(_png_bytes())
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "Report"
    worksheet.add_image(SpreadsheetImage(image_path), "B2")
    spreadsheet = tmp_path / "report.xlsx"
    workbook.save(spreadsheet)
    workbook.close()
    sink = DocumentImageSink(tmp_path / "session" / "doc_converter", source_stem="report")

    result = XlsxParser().parse(str(spreadsheet), image_sink=sink)

    assert result.visuals == ()
    assert result.warnings == ()
    assert not (tmp_path / "session" / "doc_converter").exists()


def test_xls_parser_tables() -> None:
    """XlsParser produces sheet-based Markdown tables."""
    from chrys.service.tools.builtins.doc_converter.parsers.xls import XlsParser

    mock_sheet = MagicMock()
    mock_sheet.name = "Data"
    mock_sheet.nrows = 2
    mock_sheet.ncols = 2
    mock_sheet.cell_value = lambda r, c: [["Header1", "Header2"], ["val1", "val2"]][r][c]

    mock_wb = MagicMock()
    mock_wb.sheets.return_value = [mock_sheet]

    mock_xlrd = MagicMock()
    mock_xlrd.open_workbook.return_value = mock_wb

    with patch.dict("sys.modules", {"xlrd": mock_xlrd}):
        parser = XlsParser()
        result = parser.parse("/fake.xls", image_sink=MagicMock(spec=DocumentImageSink))

    assert "# Sheet: Data" in result.markdown
    assert "| Header1 | Header2 |" in result.markdown
    assert "| val1 | val2 |" in result.markdown
    assert result.warnings == ()


def test_protocol_compliance() -> None:
    """All parsers satisfy the DocParser protocol."""
    from chrys.service.tools.builtins.doc_converter.parsers.base import DocParser
    from chrys.service.tools.builtins.doc_converter.parsers.docx import DocxParser
    from chrys.service.tools.builtins.doc_converter.parsers.pdf import PdfParser
    from chrys.service.tools.builtins.doc_converter.parsers.pptx import PptxParser
    from chrys.service.tools.builtins.doc_converter.parsers.xls import XlsParser
    from chrys.service.tools.builtins.doc_converter.parsers.xlsx import XlsxParser

    for cls in [PdfParser, DocxParser, PptxParser, XlsxParser, XlsParser]:
        assert isinstance(cls(), DocParser), f"{cls.__name__} does not satisfy DocParser protocol"


# ---------------------------------------------------------------------------
# Markdown table pipe escaping
# ---------------------------------------------------------------------------


def test_table_pipe_escaping() -> None:
    """Pipe characters in cell values are escaped so they don't break the table."""
    from chrys.service.tools.builtins.doc_converter.parsers._table import rows_to_markdown_table

    rows = [("Header", "Formula"), ("A|B", "x | y")]
    result = rows_to_markdown_table(rows)
    assert "| Header | Formula |" in result
    assert r"| A\|B | x \| y |" in result
